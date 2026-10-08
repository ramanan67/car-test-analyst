"""
app/analytics/report_generator.py

Generates PDF engineering sign-off sheets and Excel export workbooks
with cryptographic digest of test run results.
"""
from __future__ import annotations

import hashlib
import io
import json
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional


def _sha256_dict(data: Dict[str, Any]) -> str:
    """Deterministic SHA-256 of a JSON-serialisable dictionary."""
    serialised = json.dumps(data, sort_keys=True, default=str).encode()
    return hashlib.sha256(serialised).hexdigest()


class PDFReportGenerator:
    """
    Generates a PDF engineering sign-off sheet using reportlab.
    Includes a cryptographic digest of metric results for tamper-evidence.
    """

    def generate(
        self,
        run_data: Dict[str, Any],
        metrics: List[Dict[str, Any]],
        output_path: str,
    ) -> str:
        """
        Generate a PDF sign-off sheet.

        Args:
            run_data    : TestRun metadata dict.
            metrics     : List of TestResultMetric dicts.
            output_path : File system path to write the PDF.

        Returns:
            SHA-256 digest of the generated PDF content.
        """
        try:
            from reportlab.lib import colors  # type: ignore
            from reportlab.lib.pagesizes import A4  # type: ignore
            from reportlab.lib.styles import getSampleStyleSheet  # type: ignore
            from reportlab.lib.units import mm  # type: ignore
            from reportlab.platypus import (
                Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
            )  # type: ignore
        except ImportError as exc:
            raise ImportError(
                "reportlab is required for PDF generation. pip install reportlab"
            ) from exc

        doc = SimpleDocTemplate(output_path, pagesize=A4, rightMargin=15*mm, leftMargin=15*mm)
        styles = getSampleStyleSheet()
        story = []

        # Header
        story.append(Paragraph("ENGINEERING TEST RUN SIGN-OFF SHEET", styles["Title"]))
        story.append(Spacer(1, 6*mm))

        # Run summary table
        digest = _sha256_dict({"run": run_data, "metrics": metrics})
        generated_at = datetime.now(timezone.utc).isoformat()
        run_info = [
            ["Run Number",       run_data.get("run_number", "—")],
            ["Vehicle",          run_data.get("prototype_code", "—")],
            ["Category",         run_data.get("category", "—")],
            ["Facility",         run_data.get("facility", "—")],
            ["Start Time",       str(run_data.get("start_time", "—"))],
            ["End Time",         str(run_data.get("end_time", "—"))],
            ["Overall Result",   run_data.get("overall_evaluation", "—")],
            ["Report Generated", generated_at],
            ["Result Digest",    digest[:32] + "..."],
        ]
        t_run = Table(run_info, colWidths=[60*mm, 110*mm])
        t_run.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.lightgrey),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(t_run)
        story.append(Spacer(1, 8*mm))

        # Metrics table
        story.append(Paragraph("Performance Metrics Detail", styles["Heading2"]))
        story.append(Spacer(1, 3*mm))

        header = ["Criteria", "Measured", "Target", "Upper", "Lower", "Variance %", "Status"]
        rows = [header]
        for m in metrics:
            status = m.get("engineer_override_status") or m.get("status", "—")
            rows.append([
                m.get("criteria_code", "—"),
                str(m.get("measured_value", "—")),
                str(m.get("applied_target", "—")),
                str(m.get("upper_limit", "—")),
                str(m.get("lower_limit", "—")),
                str(m.get("variance_pct", "—")),
                status,
            ])

        t_metrics = Table(rows, repeatRows=1)
        t_metrics.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#003366")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f4f8")]),
        ]))
        story.append(t_metrics)
        story.append(Spacer(1, 8*mm))

        # Digest footer
        story.append(Paragraph(
            f"<b>Cryptographic Result Digest (SHA-256):</b> {digest}",
            styles["Normal"],
        ))

        doc.build(story)

        # Return SHA-256 of the generated PDF file
        with open(output_path, "rb") as fh:
            return hashlib.sha256(fh.read()).hexdigest()


class ExcelReportGenerator:
    """
    Generates a multi-sheet Excel engineering report using openpyxl.
    """

    def generate(
        self,
        run_data: Dict[str, Any],
        metrics: List[Dict[str, Any]],
        output_path: str,
    ) -> str:
        """
        Generate an Excel sign-off workbook.

        Args:
            run_data    : TestRun metadata dict.
            metrics     : List of TestResultMetric dicts.
            output_path : File system path to write the .xlsx.

        Returns:
            SHA-256 digest of the generated file content.
        """
        try:
            import openpyxl  # type: ignore
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            raise ImportError(
                "openpyxl is required for Excel generation. pip install openpyxl"
            ) from exc

        wb = openpyxl.Workbook()

        # -- Sheet 1: Run Summary --
        ws_summary = wb.active
        ws_summary.title = "Run Summary"
        header_fill = PatternFill("solid", fgColor="003366")
        header_font = Font(color="FFFFFF", bold=True)

        summary_rows = [
            ("Field", "Value"),
            ("Run Number",       run_data.get("run_number")),
            ("Vehicle Prototype", run_data.get("prototype_code")),
            ("Category",         run_data.get("category")),
            ("Facility",         run_data.get("facility")),
            ("Start Time",       str(run_data.get("start_time"))),
            ("End Time",         str(run_data.get("end_time"))),
            ("Overall Evaluation", run_data.get("overall_evaluation")),
            ("Ambient Temp (°C)", run_data.get("ambient_temp_c")),
            ("Track Condition",  run_data.get("track_condition")),
            ("Executed By",      run_data.get("executed_by")),
            ("Reviewed By",      run_data.get("reviewed_by")),
            ("Result Digest",    _sha256_dict({"run": run_data, "metrics": metrics})),
        ]

        for r_idx, (label, value) in enumerate(summary_rows, start=1):
            ws_summary.cell(r_idx, 1, label)
            ws_summary.cell(r_idx, 2, value)
            if r_idx == 1:
                for col in (1, 2):
                    cell = ws_summary.cell(r_idx, col)
                    cell.fill = header_fill
                    cell.font = header_font

        ws_summary.column_dimensions["A"].width = 25
        ws_summary.column_dimensions["B"].width = 55

        # -- Sheet 2: Metrics Detail --
        ws_metrics = wb.create_sheet("Metrics Detail")
        metric_headers = [
            "Criteria Code", "Metric Name", "Measured Value", "Target",
            "Upper Limit", "Lower Limit", "Variance %", "Status", "Override Status",
            "Override Reason",
        ]
        for c_idx, h in enumerate(metric_headers, start=1):
            cell = ws_metrics.cell(1, c_idx, h)
            cell.fill = header_fill
            cell.font = header_font

        status_colors = {
            "PASS": "C6EFCE",
            "FAIL": "FFC7CE",
            "MARGINAL_DEVIATION": "FFEB9C",
            "PENDING": "DDDDDD",
        }
        for r_idx, m in enumerate(metrics, start=2):
            status = m.get("engineer_override_status") or m.get("status", "")
            row_values = [
                m.get("criteria_code"),
                m.get("metric_name"),
                m.get("measured_value"),
                m.get("applied_target"),
                m.get("upper_limit"),
                m.get("lower_limit"),
                m.get("variance_pct"),
                m.get("status"),
                m.get("engineer_override_status"),
                m.get("override_reason"),
            ]
            fill_color = status_colors.get(status, "FFFFFF")
            fill = PatternFill("solid", fgColor=fill_color)
            for c_idx, val in enumerate(row_values, start=1):
                cell = ws_metrics.cell(r_idx, c_idx, val)
                cell.fill = fill

        for i, width in enumerate([20, 30, 15, 12, 12, 12, 12, 20, 20, 40], start=1):
            ws_metrics.column_dimensions[
                openpyxl.utils.get_column_letter(i)
            ].width = width

        wb.save(output_path)

        with open(output_path, "rb") as fh:
            return hashlib.sha256(fh.read()).hexdigest()
